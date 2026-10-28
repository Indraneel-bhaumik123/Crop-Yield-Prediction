"""
Script to build the complete, publication-grade ~20-page research paper
"Seasonal Crop Yield Prediction Using Machine Learning: A Comparative Study of KNN, Gradient Boosting, XGBoost, and MLP"
Outputs: FINAL_RESEARCH_PAPER.docx
"""

import os
import csv
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
        # Header
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Seasonal Crop Yield Prediction Using Machine Learning | Capstone Research Paper")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)
        
        # Footer
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Department of Computer Science & Engineering | Academic Final-Year Project")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

def format_paragraph(p, space_before=0, space_after=6, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    format_paragraph(h, space_before=12 if level==1 else (8 if level==2 else 6),
                     space_after=4 if level==1 else (3 if level==2 else 2),
                     align=WD_ALIGN_PARAGRAPH.LEFT)
    run = h.runs[0]
    run.font.name = "Calibri"
    if level == 1:
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 58, 138) # Deep Navy
    elif level == 2:
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 41, 59) # Slate Dark
    elif level == 3:
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(51, 65, 85)
    return h

def add_body_p(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=0, space_after=5, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_equation_box(doc, eq_text, eq_num):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell_eq = table.cell(0, 0)
    cell_num = table.cell(0, 1)
    
    cell_eq.width = Inches(5.5)
    cell_num.width = Inches(0.9)
    
    set_cell_background(cell_eq, "F8FAFC")
    set_cell_background(cell_num, "F8FAFC")
    set_cell_margins(cell_eq, top=80, bottom=80, left=150, right=100)
    set_cell_margins(cell_num, top=80, bottom=80, left=50, right=100)
    
    pe = cell_eq.paragraphs[0]
    format_paragraph(pe, space_before=2, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    re = pe.add_run(eq_text)
    re.font.name = "Cambria Math"
    re.font.size = Pt(10.5)
    re.font.italic = True
    re.font.bold = True
    re.font.color.rgb = RGBColor(15, 23, 42)
    
    pn = cell_num.paragraphs[0]
    format_paragraph(pn, space_before=2, space_after=2, align=WD_ALIGN_PARAGRAPH.RIGHT)
    rn = pn.add_run(f"({eq_num})")
    rn.font.name = "Calibri"
    rn.font.size = Pt(10)
    rn.font.bold = True
    rn.font.color.rgb = RGBColor(100, 116, 139)
    
    set_table_borders(table, color="E2E8F0", sz="4")
    
    # spacing paragraph after table
    sp = doc.add_paragraph()
    format_paragraph(sp, space_before=0, space_after=4)

def add_figure_image(doc, img_path, caption_num, caption_title, caption_text, width_inches=6.0):
    if not os.path.exists(img_path):
        print(f"Warning: Image {img_path} not found.")
        return
    p_img = doc.add_paragraph()
    format_paragraph(p_img, space_before=8, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    p_img.add_run().add_picture(img_path, width=Inches(width_inches))
    
    p_cap = doc.add_paragraph()
    format_paragraph(p_cap, space_before=2, space_after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    r_bold = p_cap.add_run(f"Figure {caption_num}. {caption_title}: ")
    r_bold.font.name = "Calibri"
    r_bold.font.size = Pt(9.5)
    r_bold.font.bold = True
    r_bold.font.color.rgb = RGBColor(30, 58, 138)
    
    r_text = p_cap.add_run(caption_text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(9.0)
    r_text.font.italic = True
    r_text.font.color.rgb = RGBColor(71, 85, 105)

def add_table_from_data(doc, table_num, table_title, headers, rows, col_widths=None):
    p_cap = doc.add_paragraph()
    format_paragraph(p_cap, space_before=10, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
    r_bold = p_cap.add_run(f"Table {table_num}. {table_title}")
    r_bold.font.name = "Calibri"
    r_bold.font.size = Pt(9.5)
    r_bold.font.bold = True
    r_bold.font.color.rgb = RGBColor(30, 58, 138)
    
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Headers
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], "1E3A8A") # Dark Navy
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        format_paragraph(p, space_before=2, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(9.0)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            
    # Rows
    for r_idx, row in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=100, right=100)
            p = row_cells[c_idx].paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 or len(val) > 20 else WD_ALIGN_PARAGRAPH.CENTER
            format_paragraph(p, space_before=2, space_after=2, align=align)
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(30, 41, 59)
                
    if col_widths and len(col_widths) == len(headers):
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)
                
    set_table_borders(table, color="CBD5E1", sz="4")
    
    sp = doc.add_paragraph()
    format_paragraph(sp, space_before=0, space_after=6)

print("Setup helper functions completed.")
