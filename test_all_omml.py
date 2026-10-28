import docx
from docx.oxml import parse_xml
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

M_NS = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'

def get_omml_eq1(): # Linear Regression: \hat{y} = \beta_0 + \sum_{j=1}^p \beta_j x_j
    return f'''<m:oMath {M_NS}>
      <m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc>
      <m:r><m:t>=</m:t></m:r>
      <m:sSub><m:e><m:r><m:t>β</m:t></m:r></m:e><m:sub><m:r><m:t>0</m:t></m:r></m:sub></m:sSub>
      <m:r><m:t>+</m:t></m:r>
      <m:nary>
        <m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr>
        <m:sub><m:r><m:t>j=1</m:t></m:r></m:sub>
        <m:sup><m:r><m:t>p</m:t></m:r></m:sup>
        <m:e>
          <m:sSub><m:e><m:r><m:t>β</m:t></m:r></m:e><m:sub><m:r><m:t>j</m:t></m:r></m:sub></m:sSub>
          <m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>j</m:t></m:r></m:sub></m:sSub>
        </m:e>
      </m:nary>
    </m:oMath>'''

def get_omml_eq2(): # Standardization: z = (x - \mu_{train}) / \sigma_{train}
    return f'''<m:oMath {M_NS}>
      <m:r><m:t>z=</m:t></m:r>
      <m:f>
        <m:num>
          <m:r><m:t>x-</m:t></m:r>
          <m:sSub><m:e><m:r><m:t>μ</m:t></m:r></m:e><m:sub><m:r><m:t>train</m:t></m:r></m:sub></m:sSub>
        </m:num>
        <m:den>
          <m:sSub><m:e><m:r><m:t>σ</m:t></m:r></m:e><m:sub><m:r><m:t>train</m:t></m:r></m:sub></m:sSub>
        </m:den>
      </m:f>
    </m:oMath>'''

def get_omml_eq3(): # Euclidean: d(x, x_i) = \sqrt{\sum_{j=1}^p (x_j - x_{i,j})^2}
    return f'''<m:oMath {M_NS}>
      <m:r><m:t>d(x,</m:t></m:r>
      <m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>
      <m:r><m:t>)=</m:t></m:r>
      <m:rad>
        <m:radPr><m:degHide m:val="on"/></m:radPr>
        <m:deg/>
        <m:e>
          <m:nary>
            <m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr>
            <m:sub><m:r><m:t>j=1</m:t></m:r></m:sub>
            <m:sup><m:r><m:t>p</m:t></m:r></m:sup>
            <m:e>
              <m:sSup>
                <m:e>
                  <m:d>
                    <m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>
                    <m:e>
                      <m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>j</m:t></m:r></m:sub></m:sSub>
                      <m:r><m:t>-</m:t></m:r>
                      <m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>i,j</m:t></m:r></m:sub></m:sSub>
                    </m:e>
                  </m:d>
                </m:e>
                <m:sup><m:r><m:t>2</m:t></m:r></m:sup>
              </m:sSup>
            </m:e>
          </m:nary>
        </m:e>
      </m:rad>
    </m:oMath>'''

def get_omml_eq4(): # Manhattan: d(x, x_i) = \sum_{j=1}^p |x_j - x_{i,j}|
    return f'''<m:oMath {M_NS}>
      <m:r><m:t>d(x,</m:t></m:r>
      <m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>
      <m:r><m:t>)=</m:t></m:r>
      <m:nary>
        <m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr>
        <m:sub><m:r><m:t>j=1</m:t></m:r></m:sub>
        <m:sup><m:r><m:t>p</m:t></m:r></m:sup>
        <m:e>
          <m:d>
            <m:dPr><m:begChr m:val="|"/><m:endChr m:val="|"/></m:dPr>
            <m:e>
              <m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>j</m:t></m:r></m:sub></m:sSub>
              <m:r><m:t>-</m:t></m:r>
              <m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>i,j</m:t></m:r></m:sub></m:sSub>
            </m:e>
          </m:d>
        </m:e>
      </m:nary>
    </m:oMath>'''

def get_omml_eq5(): # Distance-weighted KNN: \hat{y}(x) = \frac{\sum w_i y_i}{\sum w_i}, w_i = \frac{1}{d_i + \epsilon}
    return f'''<m:oMath {M_NS}>
      <m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc>
      <m:r><m:t>(x)=</m:t></m:r>
      <m:f>
        <m:num>
          <m:nary>
            <m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr>
            <m:sub><m:r><m:t>i∈</m:t></m:r><m:sSub><m:e><m:r><m:t>N</m:t></m:r></m:e><m:sub><m:r><m:t>k</m:t></m:r></m:sub></m:sSub><m:r><m:t>(x)</m:t></m:r></m:sub>
            <m:sup/>
            <m:e>
              <m:sSub><m:e><m:r><m:t>w</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>
              <m:sSub><m:e><m:r><m:t>y</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>
            </m:e>
          </m:nary>
        </m:num>
        <m:den>
          <m:nary>
            <m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr>
            <m:sub><m:r><m:t>i∈</m:t></m:r><m:sSub><m:e><m:r><m:t>N</m:t></m:r></m:e><m:sub><m:r><m:t>k</m:t></m:r></m:sub></m:sSub><m:r><m:t>(x)</m:t></m:r></m:sub>
            <m:sup/>
            <m:e>
              <m:sSub><m:e><m:r><m:t>w</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>
            </m:e>
          </m:nary>
        </m:den>
      </m:f>
      <m:r><m:t>,   </m:t></m:r>
      <m:sSub><m:e><m:r><m:t>w</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>
      <m:r><m:t>=</m:t></m:r>
      <m:f>
        <m:num><m:r><m:t>1</m:t></m:r></m:num>
        <m:den>
          <m:sSub><m:e><m:r><m:t>d</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>
          <m:r><m:t>+ϵ</m:t></m:r>
        </m:den>
      </m:f>
    </m:oMath>'''

def get_omml_eq6(): # RMSE: RMSE = \sqrt{\frac{1}{n} \sum_{i=1}^n (y_i - \hat{y}_i)^2}
    return f'''<m:oMath {M_NS}>
      <m:r><m:t>RMSE=</m:t></m:r>
      <m:rad>
        <m:radPr><m:degHide m:val="on"/></m:radPr>
        <m:deg/>
        <m:e>
          <m:f>
            <m:num><m:r><m:t>1</m:t></m:r></m:num>
            <m:den><m:r><m:t>n</m:t></m:r></m:den>
          </m:f>
          <m:nary>
            <m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr>
            <m:sub><m:r><m:t>i=1</m:t></m:r></m:sub>
            <m:sup><m:r><m:t>n</m:t></m:r></m:sup>
            <m:e>
              <m:sSup>
                <m:e>
                  <m:d>
                    <m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>
                    <m:e>
                      <m:sSub><m:e><m:r><m:t>y</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>
                      <m:r><m:t>-</m:t></m:r>
                      <m:sSub>
                        <m:e><m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc></m:e>
                        <m:sub><m:r><m:t>i</m:t></m:r></m:sub>
                      </m:sSub>
                    </m:e>
                  </m:d>
                </m:e>
                <m:sup><m:r><m:t>2</m:t></m:r></m:sup>
              </m:sSup>
            </m:e>
          </m:nary>
        </m:e>
      </m:rad>
    </m:oMath>'''

def get_omml_eq7(): # MAE: MAE = \frac{1}{n} \sum_{i=1}^n |y_i - \hat{y}_i|
    return f'''<m:oMath {M_NS}>
      <m:r><m:t>MAE=</m:t></m:r>
      <m:f>
        <m:num><m:r><m:t>1</m:t></m:r></m:num>
        <m:den><m:r><m:t>n</m:t></m:r></m:den>
      </m:f>
      <m:nary>
        <m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr>
        <m:sub><m:r><m:t>i=1</m:t></m:r></m:sub>
        <m:sup><m:r><m:t>n</m:t></m:r></m:sup>
        <m:e>
          <m:d>
            <m:dPr><m:begChr m:val="|"/><m:endChr m:val="|"/></m:dPr>
            <m:e>
              <m:sSub><m:e><m:r><m:t>y</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>
              <m:r><m:t>-</m:t></m:r>
              <m:sSub>
                <m:e><m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc></m:e>
                <m:sub><m:r><m:t>i</m:t></m:r></m:sub>
              </m:sSub>
            </m:e>
          </m:d>
        </m:e>
      </m:nary>
    </m:oMath>'''

def get_omml_eq8(): # R2: R^2 = 1 - \frac{\sum (y_i - \hat{y}_i)^2}{\sum (y_i - \bar{y})^2}
    return f'''<m:oMath {M_NS}>
      <m:sSup><m:e><m:r><m:t>R</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup>
      <m:r><m:t>=1-</m:t></m:r>
      <m:f>
        <m:num>
          <m:nary>
            <m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr>
            <m:sub><m:r><m:t>i=1</m:t></m:r></m:sub>
            <m:sup><m:r><m:t>n</m:t></m:r></m:sup>
            <m:e>
              <m:sSup>
                <m:e>
                  <m:d>
                    <m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>
                    <m:e>
                      <m:sSub><m:e><m:r><m:t>y</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>
                      <m:r><m:t>-</m:t></m:r>
                      <m:sSub><m:e><m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>
                    </m:e>
                  </m:d>
                </m:e>
                <m:sup><m:r><m:t>2</m:t></m:r></m:sup>
              </m:sSup>
            </m:e>
          </m:nary>
        </m:num>
        <m:den>
          <m:nary>
            <m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr>
            <m:sub><m:r><m:t>i=1</m:t></m:r></m:sub>
            <m:sup><m:r><m:t>n</m:t></m:r></m:sup>
            <m:e>
              <m:sSup>
                <m:e>
                  <m:d>
                    <m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>
                    <m:e>
                      <m:sSub><m:e><m:r><m:t>y</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>
                      <m:r><m:t>-</m:t></m:r>
                      <m:acc><m:accPr><m:chr m:val="̄"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc>
                    </m:e>
                  </m:d>
                </m:e>
                <m:sup><m:r><m:t>2</m:t></m:r></m:sup>
              </m:sSup>
            </m:e>
          </m:nary>
        </m:den>
      </m:f>
    </m:oMath>'''

def get_omml_eq9(): # Pearson r
    return f'''<m:oMath {M_NS}>
      <m:r><m:t>r=</m:t></m:r>
      <m:f>
        <m:num>
          <m:nary>
            <m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr>
            <m:sub><m:r><m:t>i=1</m:t></m:r></m:sub>
            <m:sup><m:r><m:t>n</m:t></m:r></m:sup>
            <m:e>
              <m:d>
                <m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>
                <m:e>
                  <m:sSub><m:e><m:r><m:t>y</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>
                  <m:r><m:t>-</m:t></m:r>
                  <m:acc><m:accPr><m:chr m:val="̄"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc>
                </m:e>
              </m:d>
              <m:d>
                <m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>
                <m:e>
                  <m:sSub><m:e><m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>
                  <m:r><m:t>-</m:t></m:r>
                  <m:acc><m:accPr><m:chr m:val="̄"/></m:accPr><m:e><m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc></m:e></m:acc>
                </m:e>
              </m:d>
            </m:e>
          </m:nary>
        </m:num>
        <m:den>
          <m:rad>
            <m:radPr><m:degHide m:val="on"/></m:radPr><m:deg/>
            <m:e>
              <m:nary>
                <m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr>
                <m:sub><m:r><m:t>i=1</m:t></m:r></m:sub><m:sup><m:r><m:t>n</m:t></m:r></m:sup>
                <m:e>
                  <m:sSup>
                    <m:e>
                      <m:d>
                        <m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>
                        <m:e>
                          <m:sSub><m:e><m:r><m:t>y</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>
                          <m:r><m:t>-</m:t></m:r>
                          <m:acc><m:accPr><m:chr m:val="̄"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc>
                        </m:e>
                      </m:d>
                    </m:e>
                    <m:sup><m:r><m:t>2</m:t></m:r></m:sup>
                  </m:sSup>
                </m:e>
              </m:nary>
            </m:e>
          </m:rad>
          <m:rad>
            <m:radPr><m:degHide m:val="on"/></m:radPr><m:deg/>
            <m:e>
              <m:nary>
                <m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr>
                <m:sub><m:r><m:t>i=1</m:t></m:r></m:sub><m:sup><m:r><m:t>n</m:t></m:r></m:sup>
                <m:e>
                  <m:sSup>
                    <m:e>
                      <m:d>
                        <m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>
                        <m:e>
                          <m:sSub><m:e><m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>
                          <m:r><m:t>-</m:t></m:r>
                          <m:acc><m:accPr><m:chr m:val="̄"/></m:accPr><m:e><m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc></m:e></m:acc>
                        </m:e>
                      </m:d>
                    </m:e>
                    <m:sup><m:r><m:t>2</m:t></m:r></m:sup>
                  </m:sSup>
                </m:e>
              </m:nary>
            </m:e>
          </m:rad>
        </m:den>
      </m:f>
    </m:oMath>'''

doc = docx.Document()
eq_funcs = [
    ("Equation 1 (Linear Regression)", get_omml_eq1(), 1),
    ("Equation 2 (Standardization)", get_omml_eq2(), 2),
    ("Equation 3 (Euclidean Distance)", get_omml_eq3(), 3),
    ("Equation 4 (Manhattan Distance)", get_omml_eq4(), 4),
    ("Equation 5 (Distance-Weighted KNN)", get_omml_eq5(), 5),
    ("Equation 6 (RMSE)", get_omml_eq6(), 6),
    ("Equation 7 (MAE)", get_omml_eq7(), 7),
    ("Equation 8 (R²)", get_omml_eq8(), 8),
    ("Equation 9 (Pearson r)", get_omml_eq9(), 9),
]

for label, omml_xml, num in eq_funcs:
    p_desc = doc.add_paragraph(label)
    # Add table for 2-column display: Equation left, Number right
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_eq = tbl.cell(0, 0)
    c_num = tbl.cell(0, 1)
    c_eq.width = Inches(5.6)
    c_num.width = Inches(0.8)
    
    p_eq = c_eq.paragraphs[0]
    p_eq.paragraph_format.space_before = Pt(4)
    p_eq.paragraph_format.space_after = Pt(4)
    p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    omml_elem = parse_xml(omml_xml)
    p_eq._p.append(omml_elem)
    
    p_n = c_num.paragraphs[0]
    p_n.paragraph_format.space_before = Pt(4)
    p_n.paragraph_format.space_after = Pt(4)
    p_n.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p_n.add_run(f"({num})")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r.font.bold = True

doc.save("test_all_omml.docx")
print("Saved test_all_omml.docx successfully!")
