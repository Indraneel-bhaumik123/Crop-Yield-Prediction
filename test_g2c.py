import docx

doc = docx.Document(r'C:\Users\indra\OneDrive\Desktop\Seasonal Crop Yield Prediction\G2C.docx')
for i, p in enumerate(doc.paragraphs):
    has_img = len(p._element.xpath('.//a:blip')) > 0 or len(p._element.xpath('.//w:drawing')) > 0
    if has_img:
        prev_t = doc.paragraphs[i-1].text[:70] if i > 0 else ""
        next_t = doc.paragraphs[i+1].text[:70] if i < len(doc.paragraphs)-1 else ""
        print(f'P{i} [IMG] | PREV: "{prev_t}" | NEXT: "{next_t}"')
