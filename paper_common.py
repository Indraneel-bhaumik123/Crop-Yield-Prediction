"""
paper_common.py
Core styling, formatting, typography, and helper functions for DOCX research paper generation.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_p(p, before=0, after=5, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align

def add_p(doc, text, bold_prefix="", before=0, after=5):
    p = doc.add_paragraph()
    format_p(p, before=before, after=after)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_h(doc, text, level):
    h = doc.add_heading(text, level=level)
    format_p(h, before=14 if level==1 else (9 if level==2 else 6),
             after=4 if level==1 else (3 if level==2 else 2),
             align=WD_ALIGN_PARAGRAPH.LEFT)
    run = h.runs[0]
    run.font.name = "Calibri"
    if level == 1:
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 58, 138) # Navy
    elif level == 2:
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42) # Slate
    elif level == 3:
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(51, 65, 85)
    return h

def add_equation(doc, eq_text, eq_num):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell_eq = table.cell(0, 0)
    cell_num = table.cell(0, 1)
    cell_eq.width = Inches(5.6)
    cell_num.width = Inches(0.8)
    
    set_cell_background(cell_eq, "F8FAFC")
    set_cell_background(cell_num, "F8FAFC")
    set_cell_margins(cell_eq, top=55, bottom=55, left=120, right=80)
    set_cell_margins(cell_num, top=55, bottom=55, left=40, right=80)
    
    pe = cell_eq.paragraphs[0]
    format_p(pe, before=1, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    re = pe.add_run(eq_text)
    re.font.name = "Cambria Math"
    re.font.size = Pt(10)
    re.font.italic = True
    re.font.bold = True
    re.font.color.rgb = RGBColor(15, 23, 42)
    
    pn = cell_num.paragraphs[0]
    format_p(pn, before=1, after=1, align=WD_ALIGN_PARAGRAPH.RIGHT)
    rn = pn.add_run(f"({eq_num})")
    rn.font.name = "Calibri"
    rn.font.size = Pt(9.5)
    rn.font.bold = True
    rn.font.color.rgb = RGBColor(100, 116, 139)
    
    set_table_borders(table, color="E2E8F0", sz="4")
    sp = doc.add_paragraph()
    format_p(sp, before=0, after=3)

# MANDATORY: Figure Caption ABOVE figure!
def add_fig(doc, img_path, num, title, desc, width=5.8):
    if not os.path.exists(img_path):
        print(f"Warning: Missing figure image: {img_path}")
        return
    # Caption ABOVE image
    p_cap = doc.add_paragraph()
    format_p(p_cap, before=10, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    rb = p_cap.add_run(f"Figure {num}. {title}")
    rb.font.name = "Calibri"
    rb.font.size = Pt(9.5)
    rb.font.bold = True
    rb.font.color.rgb = RGBColor(30, 58, 138)
    
    if desc:
        p_desc = doc.add_paragraph()
        format_p(p_desc, before=0, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
        rt = p_desc.add_run(desc)
        rt.font.name = "Calibri"
        rt.font.size = Pt(8.5)
        rt.font.italic = True
        rt.font.color.rgb = RGBColor(71, 85, 105)
        
    p_img = doc.add_paragraph()
    format_p(p_img, before=2, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    p_img.add_run().add_picture(img_path, width=Inches(width))

# Table caption ABOVE table
def add_tbl(doc, num, title, headers, rows, col_widths=None):
    p_cap = doc.add_paragraph()
    format_p(p_cap, before=10, after=3, align=WD_ALIGN_PARAGRAPH.LEFT)
    rb = p_cap.add_run(f"Table {num}. {title}")
    rb.font.name = "Calibri"
    rb.font.size = Pt(9.5)
    rb.font.bold = True
    rb.font.color.rgb = RGBColor(30, 58, 138)
    
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
        p = cell.paragraphs[0]
        format_p(p, before=1, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            
    # Rows
    for r_idx, row in enumerate(rows):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=55, bottom=55, left=75, right=75)
            p = cell.paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 or len(str(val)) > 25 else WD_ALIGN_PARAGRAPH.CENTER
            format_p(p, before=1, after=1, align=align)
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(8.0)
                r.font.color.rgb = RGBColor(30, 41, 59)
                
    if col_widths and len(col_widths) == len(headers):
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)
                
    set_table_borders(table, color="CBD5E1", sz="4")
    sp = doc.add_paragraph()
    format_p(sp, before=0, after=4)
