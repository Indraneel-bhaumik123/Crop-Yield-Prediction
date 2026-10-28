"""
build_complete_faculty_submission.py
Builds the complete manuscript: Seasonal_Crop_Yield_Prediction_Faculty_Revised.docx
and the checklist: Faculty_Red_Marks_Resolution_Checklist.docx / .txt
"""

import os
import sys
import json
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BASE_DIR = r'C:\Users\indra\OneDrive\Desktop\Seasonal Crop Yield Prediction'
PROJECT_DIR = os.path.join(BASE_DIR, 'Seasonal_Crop_Yield_Prediction_Rebuilt')
OUT_DOCX = os.path.join(BASE_DIR, 'Seasonal_Crop_Yield_Prediction_Faculty_Revised.docx')
CHECKLIST_DOCX = os.path.join(BASE_DIR, 'Faculty_Red_Marks_Resolution_Checklist.docx')
CHECKLIST_TXT = os.path.join(BASE_DIR, 'Faculty_Red_Marks_Resolution_Checklist.txt')

M_NS = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=65, bottom=65, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="none"/><w:right w:val="none"/><w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideV w:val="none"/></w:tblBorders>'
    )
    tblPr.append(borders)

def format_p(p, before=0, after=5, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align

BM_ID = 500
def get_next_bm_id():
    global BM_ID
    BM_ID += 1
    return BM_ID

def add_bookmark(paragraph, bm_name):
    bm_id = get_next_bm_id()
    bm_start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{bm_id}" w:name="{bm_name}"/>')
    bm_end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{bm_id}"/>')
    paragraph._p.insert(0, bm_start)
    paragraph._p.append(bm_end)

def create_hyperlink_run(token, anchor):
    hl = parse_xml(
        f'<w:hyperlink {nsdecls("w")} w:anchor="{anchor}" w:history="1"><w:r><w:rPr><w:rStyle w:val="Hyperlink"/><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:color w:val="0563C1"/><w:u w:val="single"/><w:sz w:val="20"/></w:rPr><w:t>{token}</w:t></w:r></w:hyperlink>'
    )
    return hl

TOKEN_REGEX = re.compile(r'(\[\d+\]|(?:Figure|Fig\.)\s*\d+|Table\s*(?:\d+|A\d+)|Equation\s*\(\d+\))')

def add_p_linked(doc, text, bold_prefix="", before=0, after=5, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    format_p(p, before=before, after=after, align=align)
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
        
    parts = TOKEN_REGEX.split(text)
    for part in parts:
        if not part:
            continue
        m = TOKEN_REGEX.match(part)
        if m:
            token = m.group(1)
            anchor = ""
            if token.startswith('['):
                num = token.strip('[]')
                anchor = f"ref_{num}"
            elif token.startswith('Figure') or token.startswith('Fig.'):
                digits = re.findall(r'\d+', token)
                if digits:
                    anchor = f"fig_{digits[0]}"
            elif token.startswith('Table'):
                digits = re.findall(r'\d+|A\d+', token)
                if digits:
                    anchor = f"tab_{digits[0]}"
            elif token.startswith('Equation'):
                digits = re.findall(r'\d+', token)
                if digits:
                    anchor = f"eq_{digits[0]}"
            
            if anchor:
                p._p.append(create_hyperlink_run(token, anchor))
            else:
                r = p.add_run(token)
                r.font.name = "Times New Roman"
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(30, 41, 59)
        else:
            r = p.add_run(part)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_heading_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    if level == 1:
        format_p(h, before=14, after=5, align=WD_ALIGN_PARAGRAPH.LEFT)
        for r in h.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = RGBColor(30, 58, 138)
    elif level == 2:
        format_p(h, before=10, after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
        for r in h.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
    elif level == 3:
        format_p(h, before=7, after=3, align=WD_ALIGN_PARAGRAPH.LEFT)
        for r in h.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.italic = True
            r.font.color.rgb = RGBColor(51, 65, 85)
    return h

def add_table_above_caption(doc, tab_num, tab_title, headers, rows, col_widths=None):
    p_cap = doc.add_paragraph()
    format_p(p_cap, before=10, after=3, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_bookmark(p_cap, f"tab_{tab_num}")
    
    r_lbl = p_cap.add_run(f"Table {tab_num}. ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(10)
    r_lbl.font.bold = True
    r_lbl.font.color.rgb = RGBColor(15, 23, 42)
    
    r_title = p_cap.add_run(tab_title)
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(10)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 41, 59)
    
    table = doc.add_table(rows=len(rows)+1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="CBD5E1", sz="4")
    
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], top=70, bottom=70, left=60, right=60)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(8.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx+1].cells
        bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=50, bottom=50, left=60, right=60)
            p = row_cells[c_idx].paragraphs[0]
            format_p(p, before=0, after=0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(8.0)
                r.font.color.rgb = RGBColor(30, 41, 59)
                
    if col_widths:
        for row in table.rows:
            for c_idx, w in enumerate(col_widths):
                if c_idx < len(row.cells):
                    row.cells[c_idx].width = Inches(w)
                    
    p_sp = doc.add_paragraph()
    format_p(p_sp, before=0, after=4)
    return table

def add_figure_above_caption(doc, img_rel_path, fig_num, fig_title, fig_desc="", width=5.8):
    full_path = os.path.join(PROJECT_DIR, img_rel_path)
    if not os.path.exists(full_path):
        full_path = os.path.join(BASE_DIR, img_rel_path)
    if not os.path.exists(full_path):
        raise FileNotFoundError(f"Missing required figure file: {full_path}")
        
    p_cap = doc.add_paragraph()
    format_p(p_cap, before=10, after=3, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_bookmark(p_cap, f"fig_{fig_num}")
    
    r_lbl = p_cap.add_run(f"Figure {fig_num}. ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(9.5)
    r_lbl.font.bold = True
    r_lbl.font.color.rgb = RGBColor(15, 23, 42)
    
    r_title = p_cap.add_run(fig_title)
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(9.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 41, 59)
    
    if fig_desc:
        r_col = p_cap.add_run(": ")
        r_col.font.name = "Times New Roman"
        r_col.font.size = Pt(9.5)
        r_col.font.bold = True
        
        r_desc = p_cap.add_run(fig_desc)
        r_desc.font.name = "Times New Roman"
        r_desc.font.size = Pt(9.5)
        r_desc.font.italic = True
        r_desc.font.color.rgb = RGBColor(71, 85, 105)
        
    p_img = doc.add_paragraph()
    format_p(p_img, before=2, after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
    p_img.add_run().add_picture(full_path, width=Inches(width))
    return p_img

def add_omml_equation(doc, omml_xml, eq_num):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    c_eq = table.cell(0, 0)
    c_num = table.cell(0, 1)
    c_eq.width = Inches(5.6)
    c_num.width = Inches(0.8)
    
    set_cell_background(c_eq, "F8FAFC")
    set_cell_background(c_num, "F8FAFC")
    set_cell_margins(c_eq, top=50, bottom=50, left=120, right=60)
    set_cell_margins(c_num, top=50, bottom=50, left=40, right=60)
    
    p_eq = c_eq.paragraphs[0]
    format_p(p_eq, before=3, after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
    omml_elem = parse_xml(omml_xml)
    p_eq._p.append(omml_elem)
    
    p_num = c_num.paragraphs[0]
    format_p(p_num, before=3, after=3, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_bookmark(p_num, f"eq_{eq_num}")
    r = p_num.add_run(f"({eq_num})")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor(71, 85, 105)
    
    set_table_borders(table, color="E2E8F0", sz="4")
    p_sp = doc.add_paragraph()
    format_p(p_sp, before=0, after=4)
    return table

# ---------------------------------------------------------------------------
# Equation XML Definitions
# ---------------------------------------------------------------------------
def eq1_lr():
    return f'''<m:oMath {M_NS}><m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc><m:r><m:t>=</m:t></m:r><m:sSub><m:e><m:r><m:t>β</m:t></m:r></m:e><m:sub><m:r><m:t>0</m:t></m:r></m:sub></m:sSub><m:r><m:t>+</m:t></m:r><m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr><m:sub><m:r><m:t>j=1</m:t></m:r></m:sub><m:sup><m:r><m:t>p</m:t></m:r></m:sup><m:e><m:sSub><m:e><m:r><m:t>β</m:t></m:r></m:e><m:sub><m:r><m:t>j</m:t></m:r></m:sub></m:sSub><m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>j</m:t></m:r></m:sub></m:sSub></m:e></m:nary></m:oMath>'''

def eq2_std():
    return f'''<m:oMath {M_NS}><m:r><m:t>z=</m:t></m:r><m:f><m:num><m:r><m:t>x-</m:t></m:r><m:sSub><m:e><m:r><m:t>μ</m:t></m:r></m:e><m:sub><m:r><m:t>train</m:t></m:r></m:sub></m:sSub></m:num><m:den><m:sSub><m:e><m:r><m:t>σ</m:t></m:r></m:e><m:sub><m:r><m:t>train</m:t></m:r></m:sub></m:sSub></m:den></m:f></m:oMath>'''

def eq3_euc():
    return f'''<m:oMath {M_NS}><m:r><m:t>d(x,</m:t></m:r><m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub><m:r><m:t>)=</m:t></m:r><m:rad><m:radPr><m:degHide m:val="on"/></m:radPr><m:deg/><m:e><m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr><m:sub><m:r><m:t>j=1</m:t></m:r></m:sub><m:sup><m:r><m:t>p</m:t></m:r></m:sup><m:e><m:sSup><m:e><m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e><m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>j</m:t></m:r></m:sub></m:sSub><m:r><m:t>-</m:t></m:r><m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>i,j</m:t></m:r></m:sub></m:sSub></m:e></m:d></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:e></m:nary></m:e></m:rad></m:oMath>'''

def eq4_man():
    return f'''<m:oMath {M_NS}><m:r><m:t>d(x,</m:t></m:r><m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub><m:r><m:t>)=</m:t></m:r><m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr><m:sub><m:r><m:t>j=1</m:t></m:r></m:sub><m:sup><m:r><m:t>p</m:t></m:r></m:sup><m:e><m:d><m:dPr><m:begChr m:val="|"/><m:endChr m:val="|"/></m:dPr><m:e><m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>j</m:t></m:r></m:sub></m:sSub><m:r><m:t>-</m:t></m:r><m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>i,j</m:t></m:r></m:sub></m:sSub></m:e></m:d></m:e></m:nary></m:oMath>'''

def eq5_knn():
    return f'''<m:oMath {M_NS}><m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc><m:r><m:t>(x)=</m:t></m:r><m:f><m:num><m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr><m:sub><m:r><m:t>i∈</m:t></m:r><m:sSub><m:e><m:r><m:t>N</m:t></m:r></m:e><m:sub><m:r><m:t>k</m:t></m:r></m:sub></m:sSub><m:r><m:t>(x)</m:t></m:r></m:sub><m:sup/><m:e><m:sSub><m:e><m:r><m:t>w</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub><m:sSub><m:e><m:r><m:t>y</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub></m:e></m:nary></m:num><m:den><m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr><m:sub><m:r><m:t>i∈</m:t></m:r><m:sSub><m:e><m:r><m:t>N</m:t></m:r></m:e><m:sub><m:r><m:t>k</m:t></m:r></m:sub></m:sSub><m:r><m:t>(x)</m:t></m:r></m:sub><m:sup/><m:e><m:sSub><m:e><m:r><m:t>w</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub></m:e></m:nary></m:den></m:f><m:r><m:t>,   </m:t></m:r><m:sSub><m:e><m:r><m:t>w</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub><m:r><m:t>=</m:t></m:r><m:f><m:num><m:r><m:t>1</m:t></m:r></m:num><m:den><m:sSub><m:e><m:r><m:t>d</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub><m:r><m:t>+ϵ</m:t></m:r></m:den></m:f></m:oMath>'''

def eq6_rmse():
    return f'''<m:oMath {M_NS}><m:r><m:t>RMSE=</m:t></m:r><m:rad><m:radPr><m:degHide m:val="on"/></m:radPr><m:deg/><m:e><m:f><m:num><m:r><m:t>1</m:t></m:r></m:num><m:den><m:r><m:t>n</m:t></m:r></m:den></m:f><m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr><m:sub><m:r><m:t>i=1</m:t></m:r></m:sub><m:sup><m:r><m:t>n</m:t></m:r></m:sup><m:e><m:sSup><m:e><m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e><m:sSub><m:e><m:r><m:t>y</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub><m:r><m:t>-</m:t></m:r><m:sSub><m:e><m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub></m:e></m:d></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:e></m:nary></m:e></m:rad></m:oMath>'''

def eq7_mae():
    return f'''<m:oMath {M_NS}><m:r><m:t>MAE=</m:t></m:r><m:f><m:num><m:r><m:t>1</m:t></m:r></m:num><m:den><m:r><m:t>n</m:t></m:r></m:den></m:f><m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr><m:sub><m:r><m:t>i=1</m:t></m:r></m:sub><m:sup><m:r><m:t>n</m:t></m:r></m:sup><m:e><m:d><m:dPr><m:begChr m:val="|"/><m:endChr m:val="|"/></m:dPr><m:e><m:sSub><m:e><m:r><m:t>y</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub><m:r><m:t>-</m:t></m:r><m:sSub><m:e><m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub></m:e></m:d></m:e></m:nary></m:oMath>'''

def eq8_r2():
    return f'''<m:oMath {M_NS}><m:sSup><m:e><m:r><m:t>R</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup><m:r><m:t>=1-</m:t></m:r><m:f><m:num><m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr><m:sub><m:r><m:t>i=1</m:t></m:r></m:sub><m:sup><m:r><m:t>n</m:t></m:r></m:sup><m:e><m:sSup><m:e><m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e><m:sSub><m:e><m:r><m:t>y</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub><m:r><m:t>-</m:t></m:r><m:sSub><m:e><m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub></m:e></m:d></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:e></m:nary></m:num><m:den><m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr><m:sub><m:r><m:t>i=1</m:t></m:r></m:sub><m:sup><m:r><m:t>n</m:t></m:r></m:sup><m:e><m:sSup><m:e><m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e><m:sSub><m:e><m:r><m:t>y</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub><m:r><m:t>-</m:t></m:r><m:acc><m:accPr><m:chr m:val="̄"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc></m:e></m:d></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:e></m:nary></m:den></m:f></m:oMath>'''

def eq9_r():
    return f'''<m:oMath {M_NS}><m:r><m:t>r=</m:t></m:r><m:f><m:num><m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr><m:sub><m:r><m:t>i=1</m:t></m:r></m:sub><m:sup><m:r><m:t>n</m:t></m:r></m:sup><m:e><m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e><m:sSub><m:e><m:r><m:t>y</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub><m:r><m:t>-</m:t></m:r><m:acc><m:accPr><m:chr m:val="̄"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc></m:e></m:d><m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e><m:sSub><m:e><m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub><m:r><m:t>-</m:t></m:r><m:acc><m:accPr><m:chr m:val="̄"/></m:accPr><m:e><m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc></m:e></m:acc></m:e></m:d></m:e></m:nary></m:num><m:den><m:rad><m:radPr><m:degHide m:val="on"/></m:radPr><m:deg/><m:e><m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr><m:sub><m:r><m:t>i=1</m:t></m:r></m:sub><m:sup><m:r><m:t>n</m:t></m:r></m:sup><m:e><m:sSup><m:e><m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e><m:sSub><m:e><m:r><m:t>y</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub><m:r><m:t>-</m:t></m:r><m:acc><m:accPr><m:chr m:val="̄"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc></m:e></m:d></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:e></m:nary></m:e></m:rad><m:rad><m:radPr><m:degHide m:val="on"/></m:radPr><m:deg/><m:e><m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr><m:sub><m:r><m:t>i=1</m:t></m:r></m:sub><m:sup><m:r><m:t>n</m:t></m:r></m:sup><m:e><m:sSup><m:e><m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e><m:sSub><m:e><m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub><m:r><m:t>-</m:t></m:r><m:acc><m:accPr><m:chr m:val="̄"/></m:accPr><m:e><m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e><m:r><m:t>y</m:t></m:r></m:e></m:acc></m:e></m:acc></m:e></m:d></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:e></m:nary></m:e></m:rad></m:den></m:f></m:oMath>'''

def eq10_penalty():
    return f'''<m:oMath {M_NS}><m:r><m:t>Error Penalty=</m:t></m:r><m:f><m:num><m:sSub><m:e><m:r><m:t>RMSE</m:t></m:r></m:e><m:sub><m:r><m:t>pooled</m:t></m:r></m:sub></m:sSub><m:r><m:t>-</m:t></m:r><m:sSub><m:e><m:r><m:t>RMSE</m:t></m:r></m:e><m:sub><m:r><m:t>specific</m:t></m:r></m:sub></m:sSub></m:num><m:den><m:sSub><m:e><m:r><m:t>RMSE</m:t></m:r></m:e><m:sub><m:r><m:t>specific</m:t></m:r></m:sub></m:sSub></m:den></m:f><m:r><m:t>×100%</m:t></m:r></m:oMath>'''

def eq11_interaction():
    return f'''<m:oMath {M_NS}><m:r><m:t>Yield=</m:t></m:r><m:sSub><m:e><m:r><m:t>β</m:t></m:r></m:e><m:sub><m:r><m:t>0</m:t></m:r></m:sub></m:sSub><m:r><m:t>+</m:t></m:r><m:sSub><m:e><m:r><m:t>β</m:t></m:r></m:e><m:sub><m:r><m:t>1</m:t></m:r></m:sub></m:sSub><m:r><m:t>Rainfall+</m:t></m:r><m:sSub><m:e><m:r><m:t>β</m:t></m:r></m:e><m:sub><m:r><m:t>2</m:t></m:r></m:sub></m:sSub><m:r><m:t>Temperature+</m:t></m:r><m:sSub><m:e><m:r><m:t>β</m:t></m:r></m:e><m:sub><m:r><m:t>3</m:t></m:r></m:sub></m:sSub><m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e><m:r><m:t>Rainfall×Temperature</m:t></m:r></m:e></m:d><m:r><m:t>+ϵ</m:t></m:r></m:oMath>'''

def eq12_density():
    return f'''<m:oMath {M_NS}><m:sSub><m:e><m:r><m:t>Error</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub><m:r><m:t>=</m:t></m:r><m:sSub><m:e><m:r><m:t>β</m:t></m:r></m:e><m:sub><m:r><m:t>0</m:t></m:r></m:sub></m:sSub><m:r><m:t>+</m:t></m:r><m:sSub><m:e><m:r><m:t>β</m:t></m:r></m:e><m:sub><m:r><m:t>1</m:t></m:r></m:sub></m:sSub><m:r><m:t>log</m:t></m:r><m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e><m:sSub><m:e><m:r><m:t>N</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub></m:e></m:d><m:r><m:t>+</m:t></m:r><m:sSub><m:e><m:r><m:t>ϵ</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub></m:oMath>'''

print("Base definitions ready.")
